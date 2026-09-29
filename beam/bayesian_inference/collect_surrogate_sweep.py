r"""Collects the surrogate + hierarchical sensor sweep and draws two figures. Read-only.

    python collect_surrogate_sweep.py              # defaults to Analysis\sensors_surrogate
    python collect_surrogate_sweep.py D:\...\Analysis\sensors_surrogate
    python collect_surrogate_sweep.py --archive D:\...\Analysis\sensors_surrogate

Each run infers the population parameters (mu_E, sd_E) of the 1000 Phase 1 specimens
from the first k sensors (option A, one instrument sigma), through a Kratos-built
response surrogate. The answer per case is ONE population, the central one, with two
cautious alternatives, exactly as in plot_population_posterior.py (figure 4):

    central        mu = 50th  percentile of the mu_E samples, sigma = 50th   of sd_E
    conservative   mu = 16th  percentile of the mu_E samples, sigma = 84th   of sd_E
    extreme        mu = 2.5th percentile of the mu_E samples, sigma = 97.5th of sd_E

Percentiles are taken on each marginal separately, so "worse" always means a lower mean
stiffness together with a wider scatter. The central population is compared with two
truths from run_info.json, bias in GPa and %:

    realised     population_truth of the drawn specimens (phase1_clean_summary.json)
    prescribed   the Phase 1 generator's N(206.9, 20.69) GPa, alpha ~ N(1.0, 0.1)

Tip identity: every surrogate fits its tip GP on the same training column with the same
seed, so the tip predictions of N01, N04 and N10 must agree exactly (difference 0.0)
and so must their tip kernels. Checked on 2001 E values over the surrogate domain.

logcE is reported but is NOT comparable across cases: the data dimension (1000
specimens x k sensors) differs. Every E is also given as alpha = E / E_ref, E_ref from
the surrogate (206.9 GPa).

Writes results.csv, surrogate_results.xlsx (sheets results = results.csv, truths,
details, surrogate, notes; needs openpyxl, skipped with a message otherwise),
surrogate_population_box.png and surrogate_joint_posterior.png.
"""
import argparse
import csv
import json
import sys
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.lines import Line2D
from matplotlib.patches import Patch

from response_surrogate import ResponseSurrogate

ARCHIVE = Path(r"D:\KratosProjects\MCMC\Analysis\sensors_surrogate")

SIGMA = 3.788e-08
GPA = 1e9
E_REF_GPA = 206.9      # fallback when the surrogate identity has no e_scale_Pa
N_TIP_CHECK = 2001
PARAMS = ("mu_E", "sd_E")
TRUTHS = ("realised", "prescribed")
TRUTH_KEY = {"mu_E": "E_mean_GPa", "sd_E": "E_sd_GPa"}
LOGCE_NOTE = ("logcE is not comparable across cases: the data dimension "
              "(specimens x k sensors) differs")

# Copied from hierarchical/plot_population_posterior.py (CASES, figure 4): the percentile
# of the mu_E samples and of the sd_E samples that define each population.
CASES = [
    ("central",      50.0, 50.0),
    ("conservative", 16.0, 84.0),
    ("extreme",       2.5, 97.5),
]
CASES_SOURCE = ("hierarchical/plot_population_posterior.py, CASES (lines 227-231) and the "
                "figure-4 loop (lines 236-244): m = np.percentile(post_g[:, 0], p_mu), "
                "s = np.percentile(post_g[:, 1], p_sd), post_g = posterior / 1e9")

# colours of plot_population_posterior.py; its truth is the realised one
MEAN_C = "#1a3a5a"
FINAL_C = "#4878a8"
CASE_C = {"central": FINAL_C, "conservative": "#c98b3a", "extreme": "#b0503f"}
REAL = "crimson"
PRESC = "#2ca02c"
EDGE = "#1f3b57"
BOX = "#5b7c99"

FIELDS = (["run", "label", "n_sensors", "stations", "sensors", "n_specimens_valid",
           "n_observations", "sigma_assumed", "e_ref_GPa"]
          + [f"{c}_{s}" for c, _, _ in CASES
             for s in ("mu_GPa", "sigma_GPa", "mu_alpha", "sigma_alpha")]
          + [f"truth_{t}_{TRUTH_KEY[p]}" for t in TRUTHS for p in PARAMS]
          + [f"central_{q}_{t}_{s}" for t in TRUTHS for q in ("mu", "sigma")
             for s in ("bias_GPa", "bias_pct")]
          + ["kratos_solves", "n_train", "n_validation", "max_error_over_sigma",
             "gp_passes_gate", "random_state", "training_hash", "tip_kernel",
             "tip_max_abs_diff_m", "tip_kernel_matches", "tip_identity",
             "n_levels", "n_likelihood_calls", "n_forward_solves", "logcE",
             "wall_time_surrogate_s", "wall_time_hierarchical_s"])

SKIPPED = []


def read_json(path):
    with open(path) as f:
        return json.load(f)


def posterior_gpa(path):
    """(mu_E, sd_E) samples in GPa from posterior.npz, or None if unavailable."""
    if not path.exists():
        return None
    with np.load(path) as z:
        if "posterior" not in z.files or "names" not in z.files:
            return None
        order = [str(n) for n in z["names"]]
        if any(p not in order for p in PARAMS):
            return None
        post_g = np.asarray(z["posterior"], dtype=float) / GPA
    return post_g[:, [order.index(p) for p in PARAMS]]


def design_cases(post_g, e_ref):
    """central / conservative / extreme (mu, sigma) exactly as plot_population_posterior.py."""
    out = {}
    for name, p_mu, p_sd in CASES:
        m = np.percentile(post_g[:, 0], p_mu)
        s = np.percentile(post_g[:, 1], p_sd)
        out[name] = {"mu_GPa": float(m), "sigma_GPa": float(s),
                     "mu_alpha": float(m / e_ref), "sigma_alpha": float(s / e_ref)}
    return out


def case_label(name, c):
    """The label box of plot_population_posterior.py figure 4."""
    return (f"{name}\n$\\mu$={c['mu_GPa']:.2f}  $\\sigma$={c['sigma_GPa']:.2f} GPa\n"
            f"$\\mu_\\alpha$={c['mu_alpha']:.4f}  $\\sigma_\\alpha$={c['sigma_alpha']:.4f}")


def collect(folder):
    info_file = folder / "run_info.json"
    summary_file = folder / "hierarchical" / "summary.json"
    npz_file = folder / "hierarchical" / "posterior.npz"
    if not info_file.exists():
        SKIPPED.append((folder.name, "no run_info.json -- the sweep never reached this case"))
        return None
    info = read_json(info_file)
    stage = info.get("stages", {}).get("hierarchical", {})
    if not summary_file.exists() or stage.get("state", "ok") != "ok":
        why = (f"hierarchical stage {stage['state']}" if stage.get("state")
               else "hierarchical stage not run")
        if not summary_file.exists():
            why += ", no hierarchical/summary.json"
        SKIPPED.append((folder.name, why))
        return None

    summary = read_json(summary_file)
    data = summary["data"]
    identity = data.get("surrogate_identity", {})
    e_ref = identity.get("e_scale_Pa", E_REF_GPA * GPA) / GPA
    post_g = posterior_gpa(npz_file)
    if post_g is None:
        SKIPPED.append((folder.name, "no usable hierarchical/posterior.npz -- central, "
                                     "conservative and extreme need the samples"))
        return None

    report_file = folder / "surrogate" / "validation_report.json"
    report = read_json(report_file) if report_file.exists() else {}
    facts = info.get("surrogate", {})
    n_tr = report.get("n_training", facts.get("n_train"))
    n_va = report.get("n_validation", facts.get("n_validation"))
    stages = info.get("stages", {})

    row = {
        "run": folder.name,
        "label": info["label"],
        "n_sensors": info["n_sensors"],
        "stations": " ".join(format(x, ".1f") for x in info["stations"]),
        "sensors": " ".join(data["sensors"]),
        "n_specimens_valid": info.get("n_specimens_valid"),
        "n_observations": data["n_observations"],
        "sigma_assumed": info["sigma_assumed"],
        "sigma_summary": data["sigma_assumed_m"],
        "e_ref_GPa": e_ref,
        "kratos_solves": None if n_tr is None or n_va is None else n_tr + n_va,
        "n_train": n_tr,
        "n_validation": n_va,
        "max_error_over_sigma": report.get("gp", {}).get(
            "max_error_over_sigma", facts.get("max_error_over_sigma")),
        "gp_passes_gate": report.get("gp_passes_gate", facts.get("gp_passes_gate")),
        "random_state": identity.get("random_state"),
        "training_hash": identity.get("training_hash"),
        "tip_kernel": (identity.get("kernels") or [None])[0],
        "tip_max_abs_diff_m": None, "tip_kernel_matches": None, "tip_identity": "n/a",
        "n_levels": summary["n_levels"],
        "n_likelihood_calls": summary["n_likelihood_calls"],
        "n_forward_solves": summary["n_forward_solves"],
        "logcE": summary["logcE"],
        "wall_time_surrogate_s": stages.get("surrogate", {}).get("wall_time_s"),
        "wall_time_hierarchical_s": stages.get("hierarchical", {}).get("wall_time_s"),
        "_truths": {t: info[f"truth_{t}"] for t in TRUTHS},
        "_post_g": post_g,
        "_model": folder / "surrogate" / "response_surrogate.joblib",
        "_tip_name": (identity.get("sensor_names") or [None])[0],
    }

    row["_cases"] = design_cases(post_g, e_ref)
    for name, c in row["_cases"].items():
        for key, value in c.items():
            row[f"{name}_{key}"] = value
    central = row["_cases"]["central"]
    for t in TRUTHS:
        for p in PARAMS:
            row[f"truth_{t}_{TRUTH_KEY[p]}"] = row["_truths"][t][TRUTH_KEY[p]]
        for q, p in (("mu", "mu_E"), ("sigma", "sd_E")):
            truth = row["_truths"][t][TRUTH_KEY[p]]
            bias = central[f"{q}_GPa"] - truth
            row[f"central_{q}_{t}_bias_GPa"] = bias
            row[f"central_{q}_{t}_bias_pct"] = 100.0 * bias / truth

    # posterior summaries of the two parameters: Excel 'details' sheet only
    for p in PARAMS:
        post = summary["posterior"][p]
        m, s = post["mean"] / GPA, post["sd"] / GPA
        lo, hi = post["ci95"][0] / GPA, post["ci95"][1] / GPA
        row.update({f"{p}_mean_GPa": m, f"{p}_sd_GPa": s,
                    f"{p}_ci_lo_GPa": lo, f"{p}_ci_hi_GPa": hi})
        for t in TRUTHS:
            truth = row["_truths"][t][TRUTH_KEY[p]]
            row[f"{p}_{t}_bias_GPa"] = m - truth
            row[f"{p}_{t}_bias_pct"] = 100.0 * (m - truth) / truth
            row[f"{p}_{t}_z"] = (m - truth) / s
            row[f"{p}_{t}_covers"] = bool(lo <= truth <= hi)
    return row


def tip_identity(runs):
    """Tip predictions of every surrogate against the smallest run's, on one E grid."""
    have = [r for r in runs if r["_model"].exists()]
    if len(have) < 2:
        return {"result": "n/a", "reason": "fewer than two surrogates on file"}
    ref = have[0]
    gp_ref = ResponseSurrogate.load(ref["_model"])
    grid = np.geomspace(gp_ref.e_min, gp_ref.e_max, N_TIP_CHECK)
    tip_ref = gp_ref.predict(grid)[:, 0]
    ref["tip_max_abs_diff_m"], ref["tip_kernel_matches"] = 0.0, True
    ok = True
    for r in have[1:]:
        gp = ResponseSurrogate.load(r["_model"])
        same_domain = (gp.e_min, gp.e_max) == (gp_ref.e_min, gp_ref.e_max)
        diff = (float(np.max(np.abs(gp.predict(grid)[:, 0] - tip_ref)))
                if same_domain else float("nan"))
        r["tip_max_abs_diff_m"] = diff
        r["tip_kernel_matches"] = (r["tip_kernel"] == ref["tip_kernel"]
                                   and r["_tip_name"] == ref["_tip_name"])
        ok &= same_domain and diff == 0.0 and r["tip_kernel_matches"]
    result = "PASS" if ok else "FAIL"
    for r in have:
        r["tip_identity"] = result
    return {"result": result, "reference": ref["label"], "n_grid": N_TIP_CHECK,
            "e_min_Pa": gp_ref.e_min, "e_max_Pa": gp_ref.e_max,
            "cases": [r["label"] for r in have]}


def save(fig, out):
    with open(str(out), "wb") as f:
        fig.savefig(f, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def alpha_axis(ax, e_ref, where="top", label=r"$\alpha = E / E_{ref}$"):
    fwd, inv = (lambda e: e / e_ref), (lambda a: a * e_ref)
    if where in ("top", "bottom"):
        sec = ax.secondary_xaxis(where, functions=(fwd, inv))
        sec.set_xlabel(label, fontsize=9)
    else:
        sec = ax.secondary_yaxis(where, functions=(fwd, inv))
        sec.set_ylabel(label, fontsize=9)
    return sec


def figure_population_box(runs, out):
    """One box per case for the central population N(mu, sigma), both truths."""
    e_ref, tr = runs[0]["e_ref_GPa"], runs[0]["_truths"]
    pos = np.arange(len(runs), dtype=float)
    central = [r["_cases"]["central"] for r in runs]
    stats = [dict(med=c["mu_GPa"],
                  q1=c["mu_GPa"] - 0.6745 * c["sigma_GPa"],
                  q3=c["mu_GPa"] + 0.6745 * c["sigma_GPa"],
                  whislo=c["mu_GPa"] - 1.96 * c["sigma_GPa"],
                  whishi=c["mu_GPa"] + 1.96 * c["sigma_GPa"], fliers=[]) for c in central]

    fig, ax = plt.subplots(figsize=(11, 2.4 + 1.2 * len(runs)))
    for t, color, ls, band in (("prescribed", PRESC, ":", 0.10),
                               ("realised", REAL, "--", 0.08)):
        m, s = tr[t]["E_mean_GPa"], tr[t]["E_sd_GPa"]
        ax.axvspan(m - s, m + s, color=color, alpha=band, zorder=0)
        ax.axvline(m, color=color, ls=ls, lw=1.8, zorder=1)
    ax.bxp(stats, positions=pos, widths=0.5, vert=False, patch_artist=True,
           showfliers=False, medianprops=dict(color="white", lw=2.0),
           whiskerprops=dict(color=EDGE, lw=1.3), capprops=dict(color=EDGE, lw=1.3),
           boxprops=dict(facecolor=FINAL_C, edgecolor=EDGE, lw=1.3, alpha=0.85))

    lo = min([s["whislo"] for s in stats]
             + [tr[t]["E_mean_GPa"] - 2 * tr[t]["E_sd_GPa"] for t in TRUTHS])
    hi = max([s["whishi"] for s in stats]
             + [tr[t]["E_mean_GPa"] + 2 * tr[t]["E_sd_GPa"] for t in TRUTHS])
    span = hi - lo
    ax.set_xlim(lo - 0.03 * span, hi + 0.42 * span)
    for y, c in zip(pos, central):
        ax.text(hi + 0.03 * span, y,
                f"E = {c['mu_GPa']:.2f} $\\pm$ {c['sigma_GPa']:.2f} GPa\n"
                f"$\\alpha$ = {c['mu_alpha']:.4f} $\\pm$ {c['sigma_alpha']:.4f}",
                va="center", ha="left", fontsize=9, color=EDGE, linespacing=1.4)

    ax.set_yticks(pos)
    ax.set_yticklabels([f"{r['n_sensors']} sensor{'s' if r['n_sensors'] > 1 else ''}"
                        for r in runs])
    ax.set_ylabel("sensor count")
    ax.set_xlabel("E  [GPa]   (box = central 50% of specimens, whiskers = 95%)")
    alpha_axis(ax, e_ref)
    ax.set_title("Central stiffness population N($\\mu$, $\\sigma$) vs. sensor count"
                 f"   |   {runs[0]['n_observations']} specimens, sigma {SIGMA:.4e}",
                 fontsize=12, pad=34)
    ax.grid(axis="x", alpha=0.25)
    ax.legend(handles=[
        Patch(facecolor=FINAL_C, edgecolor=EDGE, alpha=0.85,
              label="central population (50th percentiles of $\\mu_E$, $\\sigma_E$)"),
        Line2D([], [], color=REAL, ls="--", lw=1.8,
               label=f"realised {tr['realised']['E_mean_GPa']:.2f} $\\pm$ "
                     f"{tr['realised']['E_sd_GPa']:.2f} GPa"),
        Line2D([], [], color=PRESC, ls=":", lw=1.8,
               label=f"prescribed {tr['prescribed']['E_mean_GPa']:.2f} $\\pm$ "
                     f"{tr['prescribed']['E_sd_GPa']:.2f} GPa"),
        Patch(facecolor=REAL, alpha=0.20, label="realised mean $\\pm$ 1 SD"),
        Patch(facecolor=PRESC, alpha=0.25, label="prescribed mean $\\pm$ 1 SD"),
    ], loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=3, fontsize=9, frameon=False)
    save(fig, out)


def figure_joint_posterior(runs, out):
    """One panel per case: posterior (mu_E, sd_E) pairs, the three design populations
    with their figure-4 label boxes, and both truths. Shared axis ranges."""
    e_ref, tr = runs[0]["e_ref_GPa"], runs[0]["_truths"]
    xs = np.concatenate([r["_post_g"][:, 0] for r in runs]
                        + [[tr[t]["E_mean_GPa"] for t in TRUTHS]])
    ys = np.concatenate([r["_post_g"][:, 1] for r in runs]
                        + [[tr[t]["E_sd_GPa"] for t in TRUTHS]])
    pad_x, pad_y = 0.06 * np.ptp(xs), 0.06 * np.ptp(ys)

    fig, axes = plt.subplots(len(runs), 1, figsize=(11, 3.9 * len(runs)),
                             sharex=True, sharey=True, squeeze=False)
    fig.subplots_adjust(right=0.66, hspace=0.18)
    for ax, r in zip(axes[:, 0], runs):
        post_g = r["_post_g"]
        ax.plot(post_g[:, 0], post_g[:, 1], ".", ms=2.5, alpha=0.35, color=FINAL_C)
        ax.plot(tr["realised"]["E_mean_GPa"], tr["realised"]["E_sd_GPa"], "x",
                color=REAL, ms=11, mew=2.2, zorder=6)
        ax.plot(tr["prescribed"]["E_mean_GPa"], tr["prescribed"]["E_sd_GPa"], "+",
                color=PRESC, ms=13, mew=2.2, zorder=6)
        for i, (name, _, _) in enumerate(CASES):
            c = r["_cases"][name]
            ax.plot(c["mu_GPa"], c["sigma_GPa"], "o", ms=8, mfc=CASE_C[name], mec=MEAN_C,
                    mew=1.4, zorder=7)
            ax.annotate(case_label(name, c), xy=(c["mu_GPa"], c["sigma_GPa"]),
                        xytext=(1.03, 0.84 - 0.34 * i), textcoords="axes fraction",
                        va="center", ha="left", fontsize=8, color=MEAN_C, linespacing=1.3,
                        bbox=dict(boxstyle="round,pad=0.35", fc="white", ec=CASE_C[name],
                                  lw=1.3),
                        arrowprops=dict(arrowstyle="-", color=CASE_C[name], lw=0.9))
        r_corr = np.corrcoef(post_g[:, 0], post_g[:, 1])[0, 1]
        ax.set_title(f"{r['label']}: {r['n_sensors']} sensor"
                     f"{'s' if r['n_sensors'] > 1 else ''}   corr = {r_corr:+.3f}",
                     fontsize=11, loc="left")
        ax.set_ylabel("$\\sigma_E$ [GPa]")
        ax.grid(alpha=0.25)
    ax = axes[0, 0]
    ax.set_xlim(xs.min() - pad_x, xs.max() + pad_x)
    ax.set_ylim(ys.min() - pad_y, ys.max() + pad_y)
    alpha_axis(ax, e_ref, "top", "$\\mu_\\alpha$")
    axes[-1, 0].set_xlabel("$\\mu_E$ [GPa]")
    axes[0, 0].legend(handles=[
        Line2D([], [], marker=".", ls="", color=FINAL_C, ms=6, label="posterior samples"),
        Line2D([], [], marker="x", ls="", color=REAL, ms=9, mew=2.0,
               label=f"realised truth ({tr['realised']['E_mean_GPa']:.2f}, "
                     f"{tr['realised']['E_sd_GPa']:.2f})"),
        Line2D([], [], marker="+", ls="", color=PRESC, ms=11, mew=2.0,
               label=f"prescribed truth ({tr['prescribed']['E_mean_GPa']:.2f}, "
                     f"{tr['prescribed']['E_sd_GPa']:.2f})"),
    ], loc="upper right", fontsize=8, framealpha=0.9)
    fig.suptitle("Joint posterior of ($\\mu_E$, $\\sigma_E$) and the three design "
                 "populations per sensor count", fontsize=12, x=0.4, y=1.0)
    save(fig, out)


NUM2, INT, SCI, ALPHA = "0.00", "0", "0.000E+00", "0.0000"


def raw_format(name):
    """Number format of a results.csv column in the results sheet."""
    if "GPa" in name or name.endswith("_pct") or name == "logcE":
        return NUM2
    if name.endswith("_alpha"):
        return ALPHA
    if name in ("sigma_assumed", "max_error_over_sigma", "tip_max_abs_diff_m"):
        return SCI
    if name in ("n_sensors", "n_specimens_valid", "n_observations", "kratos_solves",
                "n_train", "n_validation", "random_state", "n_levels",
                "n_likelihood_calls", "n_forward_solves", "wall_time_surrogate_s",
                "wall_time_hierarchical_s"):
        return INT
    return None


def shown(value, fmt):
    """Roughly what Excel displays, for the column width."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "FALSE"
    if isinstance(value, float) and fmt and "E" not in fmt:
        return f"{value:.{len(fmt.split('.')[1]) if '.' in fmt else 0}f}"
    if isinstance(value, float) and fmt:
        return f"{value:.3e}"
    return str(value)


def fill_sheet(ws, headers, rows, formats, font):
    """Bold frozen header, one number format per column, widths from the content."""
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = font
    ws.freeze_panes = "A2"
    for j, (header, fmt) in enumerate(zip(headers, formats), 1):
        column = [row[j - 1] for row in rows]
        if fmt:
            for i, value in enumerate(column, 2):
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    ws.cell(i, j).number_format = fmt
        width = max([len(str(header))] + [len(shown(v, fmt)) for v in column])
        ws.column_dimensions[ws.cell(1, j).column_letter].width = min(width + 2, 60)


def write_excel(runs, csv_rows, tip, archive, path):
    """surrogate_results.xlsx; returns True if written."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print(f"openpyxl is not installed -- {path.name} skipped "
              f"(install it with: {sys.executable} -m pip install openpyxl)")
        return False

    bold = Font(bold=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "results"
    fill_sheet(ws, FIELDS, [[row[k] for k in FIELDS] for row in csv_rows],
               [raw_format(k) for k in FIELDS], bold)

    ws = wb.create_sheet("truths")
    headers = ["run", "sensors", "quantity", "truth", "truth [GPa]", "central [GPa]",
               "bias [GPa]", "bias %"]
    rows = [[r["label"], r["n_sensors"], q, t, r[f"truth_{t}_{TRUTH_KEY[p]}"],
             r[f"central_{q}_GPa"], r[f"central_{q}_{t}_bias_GPa"],
             r[f"central_{q}_{t}_bias_pct"]]
            for r in runs for q, p in (("mu", "mu_E"), ("sigma", "sd_E")) for t in TRUTHS]
    fill_sheet(ws, headers, rows, [None, INT, None, None, NUM2, NUM2, NUM2, NUM2], bold)

    ws = wb.create_sheet("details")
    headers = ["run", "sensors", "parameter", "truth", "truth [GPa]", "posterior mean [GPa]",
               "posterior SD [GPa]", "95% lo [GPa]", "95% hi [GPa]", "bias [GPa]",
               "bias %", "z", "inside 95% CI"]
    rows = [[r["label"], r["n_sensors"], p, t, r[f"truth_{t}_{TRUTH_KEY[p]}"],
             r[f"{p}_mean_GPa"], r[f"{p}_sd_GPa"], r[f"{p}_ci_lo_GPa"], r[f"{p}_ci_hi_GPa"],
             r[f"{p}_{t}_bias_GPa"], r[f"{p}_{t}_bias_pct"], r[f"{p}_{t}_z"],
             r[f"{p}_{t}_covers"]]
            for r in runs for p in PARAMS for t in TRUTHS]
    fill_sheet(ws, headers, rows, [None, INT, None, None] + [NUM2] * 8 + [None], bold)

    ws = wb.create_sheet("surrogate")
    headers = ["run", "sensors", "Kratos solves", "n train", "n validation",
               "max err / sigma", "gate pass", "random_state", "training hash", "tip kernel",
               "tip max |diff| vs ref [m]", "tip kernel matches", "tip identity",
               "levels", "likelihood calls", "forward solves (hierarchical)",
               "logcE (not comparable across cases)", "wall surrogate [s]",
               "wall hierarchical [s]"]
    rows = [[r["label"], r["n_sensors"], r["kratos_solves"], r["n_train"], r["n_validation"],
             r["max_error_over_sigma"], r["gp_passes_gate"], r["random_state"],
             r["training_hash"], r["tip_kernel"], r["tip_max_abs_diff_m"],
             r["tip_kernel_matches"], r["tip_identity"], r["n_levels"],
             r["n_likelihood_calls"], r["n_forward_solves"], r["logcE"],
             r["wall_time_surrogate_s"], r["wall_time_hierarchical_s"]] for r in runs]
    fill_sheet(ws, headers, rows, [None, INT, INT, INT, INT, SCI, None, INT, None, None,
                                   SCI, None, None, INT, INT, INT, NUM2, INT, INT], bold)

    ws = wb.create_sheet("notes")
    tr = runs[0]["_truths"]
    rows = [["archive", str(archive)],
            ["collected", datetime.now().isoformat(timespec="seconds")]]
    rows += [[f"{name}", f"mu = {p_mu:g}th percentile of the posterior mu_E samples, "
                         f"sigma = {p_sd:g}th percentile of the posterior sd_E samples"]
             for name, p_mu, p_sd in CASES]
    rows += [["percentiles", "taken on each marginal separately, so a lower mean always "
                             "comes with a wider scatter; not a single joint posterior draw"],
             ["definition source", CASES_SOURCE],
             ["logcE", LOGCE_NOTE],
             ["sigma per sensor [m]", ", ".join(f"{s:.4e}" for s in
                                                sorted({r["sigma_assumed"] for r in runs}))],
             ["E_ref [GPa]", f"{runs[0]['e_ref_GPa']:g}"],
             ["truth realised", f"E {tr['realised']['E_mean_GPa']:.4f} +- "
                                f"{tr['realised']['E_sd_GPa']:.4f} GPa; "
                                f"source {tr['realised'].get('source', '')}"],
             ["truth prescribed", f"E {tr['prescribed']['E_mean_GPa']:.4f} +- "
                                  f"{tr['prescribed']['E_sd_GPa']:.4f} GPa; "
                                  f"source {tr['prescribed'].get('source', '')}"],
             ["tip identity", f"{tip['result']}"
                              + (f" (reference {tip['reference']}, {tip['n_grid']} E values, "
                                 f"cases {', '.join(tip['cases'])})" if "reference" in tip
                                 else f" ({tip.get('reason', '')})")],
             ["details sheet", "posterior mean / SD / 95% CI of mu_E and sd_E; "
                               "z = bias / posterior SD"]]
    rows += [[f"sensors {r['label']}", r["sensors"]] for r in runs]
    fill_sheet(ws, ["item", "value"], rows, [None, None], bold)

    try:
        wb.save(path)
    except PermissionError:
        print(f"close {path} (probably open in Excel) and run again")
        return False
    return True


def main():
    ap = argparse.ArgumentParser(description="collect the surrogate sensor sweep")
    ap.add_argument("archive", nargs="?", type=Path, default=None)
    ap.add_argument("--archive", dest="archive_opt", type=Path, default=None)
    args = ap.parse_args()
    archive = args.archive_opt or args.archive or ARCHIVE
    if not archive.exists():
        sys.exit(f"no such archive: {archive}")

    folders = [f for f in sorted(archive.iterdir()) if f.is_dir()]
    runs = [r for r in (collect(f) for f in folders) if r]
    if SKIPPED:
        print(f"*** {len(SKIPPED)} of {len(folders)} run folders have no result "
              f"and are NOT in the tables below ***")
        for name, why in SKIPPED:
            print(f"    {name:<20} {why}")
        print()
    if not runs:
        sys.exit(f"no usable hierarchical results found under {archive}")
    runs.sort(key=lambda r: r["n_sensors"])

    sig = {r["sigma_assumed"] for r in runs} | {r["sigma_summary"] for r in runs}
    if len(sig) > 1:
        print("WARNING: sigma varies across runs or between run_info and summary -- "
              "this sweep should hold it fixed")
    elif abs(sig.pop() - SIGMA) / SIGMA > 1e-3:
        print(f"WARNING: sigma_assumed is not {SIGMA:.4e}")
    for t in TRUTHS:
        if len({json.dumps(r["_truths"][t], sort_keys=True) for r in runs}) > 1:
            print(f"WARNING: the {t} truth differs between runs -- mixed datasets?")
    if len({r["n_observations"] for r in runs}) > 1:
        print("WARNING: the number of specimens differs between runs")
    if runs[0]["n_sensors"] != 1:
        print("WARNING: no N=1 run -- the tip identity is checked against the smallest "
              "run present")

    tip = tip_identity(runs)
    csv_rows = [{k: r[k] for k in FIELDS} for r in runs]
    out = archive / "results.csv"
    xlsx = archive / "surrogate_results.xlsx"
    written = []
    try:
        with open(out, "w", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(csv_rows)
        written.append(out)
    except PermissionError:
        print(f"close {out} (probably open in Excel) and run again")
    if write_excel(runs, csv_rows, tip, archive, xlsx):
        written.append(xlsx)

    tr, e_ref = runs[0]["_truths"], runs[0]["e_ref_GPa"]
    print(f"E_ref = {e_ref:g} GPa")
    for t in TRUTHS:
        print(f"truth {t:<10}: E {tr[t]['E_mean_GPa']:8.4f} +- {tr[t]['E_sd_GPa']:7.4f} GPa"
              f"   alpha {tr[t]['E_mean_GPa'] / e_ref:.6f} +- {tr[t]['E_sd_GPa'] / e_ref:.6f}")
    print()

    header = (f"{'run':<6}{'k':>4}{'population':>14}{'E mean':>9}{'E sd':>8}"
              f"{'alpha':>9}{'a sd':>8}{'mu real':>9}{'%':>7}{'sd real':>9}{'%':>7}"
              f"{'mu presc':>10}{'%':>7}{'sd presc':>10}{'%':>7}")
    print("population per case   (central = answer; bias columns [GPa, %] for central only)")
    print(header)
    print("-" * len(header))
    for r in runs:
        for name, _, _ in CASES:
            c = r["_cases"][name]
            line = (f"{r['label'] if name == 'central' else '':<6}"
                    f"{r['n_sensors'] if name == 'central' else '':>4}{name:>14}"
                    f"{c['mu_GPa']:>9.2f}{c['sigma_GPa']:>8.2f}"
                    f"{c['mu_alpha']:>9.4f}{c['sigma_alpha']:>8.4f}")
            if name == "central":
                for t in TRUTHS:
                    width = 10 if t == "prescribed" else 9     # matches the header
                    for q in ("mu", "sigma"):
                        line += (f"{r[f'central_{q}_{t}_bias_GPa']:>+{width}.2f}"
                                 f"{r[f'central_{q}_{t}_bias_pct']:>+7.2f}")
            print(line)

    def num(v, spec):
        return "" if v is None else format(v, spec)

    header = (f"{'run':<6}{'k':>4}{'solves':>8}{'train':>7}{'valid':>7}{'err/sig':>11}"
              f"{'gate':>6}{'seed':>10}{'levels':>8}{'calls':>8}{'logcE*':>12}"
              f"{'t_surr':>8}{'t_hier':>8}")
    print("\nsurrogate and hierarchical   (* " + LOGCE_NOTE + ")")
    print(header)
    print("-" * len(header))
    for r in runs:
        gate = ("" if r["gp_passes_gate"] is None
                else "PASS" if r["gp_passes_gate"] else "FAIL")
        print(f"{r['label']:<6}{r['n_sensors']:>4}{num(r['kratos_solves'], '>8'):>8}"
              f"{num(r['n_train'], '>7'):>7}{num(r['n_validation'], '>7'):>7}"
              f"{num(r['max_error_over_sigma'], '>11.3e'):>11}{gate:>6}"
              f"{num(r['random_state'], '>10'):>10}{r['n_levels']:>8}"
              f"{r['n_likelihood_calls']:>8}{r['logcE']:>12.2f}"
              f"{num(r['wall_time_surrogate_s'], '>8.0f'):>8}"
              f"{num(r['wall_time_hierarchical_s'], '>8.0f'):>8}")

    print(f"\ntip identity: {tip['result']}", end="")
    if "reference" in tip:
        print(f"   (tip GP vs {tip['reference']}, {tip['n_grid']} E values over "
              f"{tip['e_min_Pa'] / GPA:g}-{tip['e_max_Pa'] / GPA:g} GPa)")
        for r in runs:
            if r["tip_max_abs_diff_m"] is not None:
                print(f"  {r['label']:<6} max |tip difference| = {r['tip_max_abs_diff_m']!r} m"
                      f"   kernel {'matches' if r['tip_kernel_matches'] else 'DIFFERS'}: "
                      f"{r['tip_kernel']}")
    else:
        print(f"   ({tip['reason']})")

    print()
    for r in runs:
        c = r["_cases"]["central"]
        print(f"{r['label']}: E = {c['mu_GPa']:.2f} +- {c['sigma_GPa']:.2f} GPa  "
              f"(alpha {c['mu_alpha']:.4f} +- {c['sigma_alpha']:.4f};  realised "
              f"{tr['realised']['E_mean_GPa']:.2f} +- {tr['realised']['E_sd_GPa']:.2f}, "
              f"prescribed {tr['prescribed']['E_mean_GPa']:.2f} +- "
              f"{tr['prescribed']['E_sd_GPa']:.2f})")

    figures = [("surrogate_population_box.png", figure_population_box),
               ("surrogate_joint_posterior.png", figure_joint_posterior)]
    for name, draw in figures:
        draw(runs, archive / name)
        written.append(archive / name)
    print("\nwrote " + "\n      ".join(str(p) for p in written))


if __name__ == "__main__":
    main()
